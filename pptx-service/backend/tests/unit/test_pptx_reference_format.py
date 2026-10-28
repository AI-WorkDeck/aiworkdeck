# SPDX-FileCopyrightText: 2026 北京京微资易科技有限公司 and AI WorkDeck contributors
# SPDX-License-Identifier: AGPL-3.0-or-later
"""Reference inspection uses actual saved PPTX files without AI/service dependencies."""
import importlib.util
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.util import Inches, Pt

_spec = importlib.util.spec_from_file_location(
    'pptx_reference_inspection', Path(__file__).parents[2] / 'services/pptx_format_service.py')
pfs = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(pfs)


def test_reference_geometry_background_theme_and_direct_paragraph_format(tmp_path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string('123456')
    shape = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1))
    para = shape.text_frame.paragraphs[0]
    para.space_before, para.space_after = Pt(6), Pt(9)
    para._p.get_or_add_pPr().set('marL', str(int(Pt(18))))
    para._p.get_or_add_pPr().set('indent', str(int(Pt(-9))))
    run = para.add_run()
    run.text = 'Synthetic reference'
    run.font.size, run.font.bold = Pt(24), True
    run.font.color.theme_color = MSO_THEME_COLOR.ACCENT_1
    path = tmp_path / 'reference.pptx'
    prs.save(path)
    data = pfs.inspect(str(path))
    assert data['page_size_pt']['height'] == 540
    page = data['slides'][0]
    assert page['background']['source'] == 'slide'
    assert page['background']['fill']['color']['value'] == '123456'
    assert page['theme']['colors']['accent1']['value']
    assert page['layout_name'] == 'Blank'
    item = page['shapes'][0]
    assert item['bounds_pt'] == {'left': 72, 'top': 144, 'width': 288, 'height': 72}
    p = item['paragraphs'][0]
    assert (p['space_before_pt'], p['space_after_pt']) == (6, 9)
    assert (p['left_margin_pt'], p['first_line_indent_pt']) == (18, -9)
    assert p['runs'][0]['color_reference']['value'] == 'accent1'
    assert p['runs'][0]['color'] is None  # Unresolved theme is not fabricated RGB.
    assert p['runs'][0]['size_pt'] == 24


def test_inherited_background_and_font_are_not_reported_as_missing(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = 'Inherited title'
    slide.slide_layout.background.fill.solid()
    slide.slide_layout.background.fill.fore_color.rgb = RGBColor.from_string('DDEEFF')
    path = tmp_path / 'inherited.pptx'
    prs.save(path)
    data = pfs.inspect(str(path))
    page = data['slides'][0]
    assert page['background']['source'] == 'layout'
    assert page['background']['fill']['color']['value'] == 'DDEEFF'
    assert page['shapes'][0]['paragraphs'][0]['runs'][0]['size_pt'] is None
    assert page['master_text_styles_xml']
    assert page['layout_placeholders']
    assert page['master_placeholders']
    assert 'inherited/unspecified' in data['format_semantics']
    assert page['theme']['font_scheme_xml']


def test_group_children_and_table_dimensions_are_visible(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(Pt(10), Pt(20), Pt(80), Pt(30)).text = 'Grouped text'
    slide.shapes.add_table(2, 2, Pt(50), Pt(80), Pt(300), Pt(100))
    path = tmp_path / 'group.pptx'
    prs.save(path)
    page = pfs.inspect(str(path))['slides'][0]
    assert page['shapes'][0]['children'][0]['paragraphs'][0]['text'] == 'Grouped text'
    assert page['shapes'][0]['group_transform_xml']
    assert page['shapes'][1]['table']['column_widths_pt'] == [150, 150]


def test_missing_optional_transform_preserves_text_and_unknown_geometry(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_textbox(Pt(10), Pt(20), Pt(80), Pt(30))
    shape.text = 'Geometry unspecified'
    props = shape._element.spPr
    props.remove(props.xfrm)
    path = tmp_path / 'missing-transform.pptx'
    prs.save(path)
    item = pfs.inspect(str(path))['slides'][0]['shapes'][0]
    assert item['paragraphs'][0]['text'] == 'Geometry unspecified'
    assert item['bounds_pt'] == dict.fromkeys(('left', 'top', 'width', 'height'))
