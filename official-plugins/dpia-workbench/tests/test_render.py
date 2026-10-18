# SPDX-FileCopyrightText: 2026 北京京微资易科技有限公司 and AI WorkDeck contributors
# SPDX-License-Identifier: AGPL-3.0-or-later
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
import zipfile
from docx import Document
from pypdf import PdfReader

RUNTIME = Path(__file__).resolve().parents[1] / 'runtime'


class ReportRenderingTest(unittest.TestCase):
    def test_real_word_pdf_pages_preserve_original_and_native_comment(self):
        with tempfile.TemporaryDirectory(prefix='awd-dpia-report-') as folder:
            root = Path(folder)
            document = Document()
            document.add_heading('DPIA 合成报告 / Synthetic report', 0)
            run = document.add_paragraph().add_run('保存期限尚待确认。Retention period is unknown.')
            document.add_comment(run, text='合成测试：需业务方提供依据，不代填事实。', author='QA')
            document.add_page_break()
            document.add_paragraph('第二页：本文件仅用于渲染验证，不构成法律结论。')
            source = root / 'report.docx'
            document.save(source)
            original = source.read_bytes()
            result = subprocess.run([sys.executable, str(RUNTIME / 'render.py'), str(source),
                                     '--output_dir', str(root / 'render'), '--emit_pdf'],
                                    check=True, capture_output=True, text=True, timeout=240)
            receipt = json.loads(result.stdout)
            self.assertTrue(Path(receipt['pdf']).is_file())
            pages = len(PdfReader(receipt['pdf']).pages)
            self.assertGreaterEqual(pages, 2)
            for index in range(1, pages + 1):
                self.assertGreater((root / 'render' / f'page-{index}.png').stat().st_size, 1000)
            self.assertEqual(source.read_bytes(), original)
            with zipfile.ZipFile(source) as archive:
                self.assertIn('word/comments.xml', archive.namelist())


if __name__ == '__main__':
    unittest.main()
