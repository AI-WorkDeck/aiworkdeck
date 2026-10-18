"""Export the already-adopted internal report to Word. No reasoning or formal approval."""
import json, sys
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def write(args):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2.2)
    for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3', 'List Bullet', 'List Number']:
        style = doc.styles[name]
        style.font.name = 'Calibri'
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'PingFang SC')
    normal = doc.styles['Normal']
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.2
    doc.add_paragraph(args['title'] + ' 内部评估报告', 'Title')
    doc.add_paragraph('本文件依据已接收的分析记录生成，供内部审阅。保留原报告中的条件、未知与待解决事项，不代表正式交付、风险接受或批准上线。')
    doc.add_paragraph('依据材料版本 ' + str(args['inputVersion']) + '；分析版本 ' + args['revisionId'])
    for block in args['blocks']:
        kind = block['type']
        if kind == 'heading':
            doc.add_paragraph(block['text'], 'Heading ' + str(min(3, block['level'])))
        elif kind == 'list':
            for item in block['items']:
                doc.add_paragraph(item, 'List Number' if block['ordered'] else 'List Bullet')
        elif kind == 'table':
            rows = block['rows']
            # Wide matrices read better as labeled records; all cells are retained.
            if len(rows[0]) > 4:
                for row in rows[1:]:
                    for index, value in enumerate(row):
                        doc.add_paragraph(rows[0][index] + '：' + value)
                if len(rows) == 1:
                    doc.add_paragraph('；'.join(rows[0]))
                continue
            table = doc.add_table(rows=1, cols=len(rows[0]))
            table.style = 'Table Grid'
            for i, text in enumerate(rows[0]):
                table.rows[0].cells[i].text = text
                for run in table.rows[0].cells[i].paragraphs[0].runs: run.bold = True
                shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'EDF1F5'); table.rows[0].cells[i]._tc.get_or_add_tcPr().append(shade)
            repeat = OxmlElement('w:tblHeader'); table.rows[0]._tr.get_or_add_trPr().append(repeat)
            for row in rows[1:]:
                cells = table.add_row().cells
                for i, text in enumerate(row): cells[i].text = text
            borders = OxmlElement('w:tblBorders')
            for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
                node = OxmlElement('w:' + edge)
                for key, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]: node.set(qn('w:' + key), value)
                borders.append(node)
            table._tbl.tblPr.append(borders)
            for row in table.rows:
                for cell in row.cells:
                    tcpr = cell._tc.get_or_add_tcPr()
                    margins = OxmlElement('w:tcMar')
                    for edge in ['top', 'bottom', 'left', 'right']:
                        node = OxmlElement('w:' + edge); node.set(qn('w:w'), '100'); node.set(qn('w:type'), 'dxa'); margins.append(node)
                    tcpr.append(margins)
                    cell.vertical_alignment = 1
            doc.add_paragraph()
        else:
            doc.add_paragraph(block['text'])
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run('内部审阅  第 ')
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    footer.add_run(' 页')
    path = Path(args['out'])
    with path.open('xb') as stream: doc.save(stream)
    # Verify the saved archive opens before publishing the receipt.
    Document(path)
    return {'document': str(path)}

if __name__ == '__main__':
    try: print(json.dumps(write(json.load(sys.stdin)), ensure_ascii=False))
    except Exception as exc:
        print(json.dumps({'error': str(exc)[:1000]})); sys.exit(1)
