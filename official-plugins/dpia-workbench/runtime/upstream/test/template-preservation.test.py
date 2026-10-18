import unittest,importlib.util,tempfile,json,hashlib,subprocess,sys
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches
from docx.enum.text import WD_BREAK
ROOT=Path(__file__).resolve().parents[1]
spec=importlib.util.spec_from_file_location('composer',ROOT/'report-compose.py');m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m)
class TemplatePreservation(unittest.TestCase):
 def test_empty_page_break_moves_to_heading_without_losing_content(self):
  d=Document();d.add_paragraph('Previous section');d.add_paragraph();d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
  h=d.add_heading('Next section',1);d.add_paragraph('Answer')
  m.normalize_template_pagination(d,set())
  self.assertEqual([p.text for p in d.paragraphs],['Previous section','Next section','Answer'])
  self.assertTrue(h.paragraph_format.page_break_before)
  self.assertEqual(len(list(d.element.body.iter(qn('w:br')))),0)
 def test_semantic_and_double_page_breaks_are_preserved(self):
  for kind in ['text','bookmark','double','section']:
   with self.subTest(kind=kind):
    d=Document();p=d.add_paragraph('Retained' if kind=='text' else '')
    p.add_run().add_break(WD_BREAK.PAGE)
    if kind=='bookmark':
     b=m.OxmlElement('w:bookmarkStart');b.set(qn('w:id'),'8');b.set(qn('w:name'),'kept');p._p.append(b)
    if kind=='double':p.add_run().add_break(WD_BREAK.PAGE)
    if kind=='section':p._p.get_or_add_pPr().append(m.OxmlElement('w:sectPr'))
    d.add_heading('Next',1);before=d.element.body.xml
    m.normalize_template_pagination(d,set());self.assertEqual(d.element.body.xml,before)
 def test_real_template_page_boundaries_preserve_words_and_bookmarks(self):
  d=Document(ROOT/'skill/assets/official/edpb-2026-v1-en.docx')
  text=lambda:''.join(t.text or '' for t in d.element.body.iter(qn('w:t')))
  bookmarks=lambda:[dict(b.attrib) for b in d.element.body.iter(qn('w:bookmarkStart'))]
  before=text();anchors=bookmarks();m.normalize_template_pagination(d,set())
  h=next(p for p in d.paragraphs if p.text.strip()=='RISK ASSESSMENT AND MANAGEMENT')
  self.assertTrue(h.paragraph_format.page_break_before);self.assertEqual(before,text());self.assertEqual(anchors,bookmarks())
  self.assertNotEqual(h._p.getprevious().tag,qn('w:p'))
 def test_reference_row_cannot_split_but_long_row_can(self):
  d,review,answer=self.composed();row=d.tables[0].rows[1]._tr
  self.assertIsNotNone(row.find(qn('w:trPr')).find(qn('w:cantSplit')))
  self.assertEqual(self.answer_errors(),[])
  d=Document();t=d.add_table(rows=2,cols=2);t.cell(0,0).text='F001';t.cell(0,1).text='Long answer. '*500
  t.cell(1,0).text='F002';t.cell(1,1).text='Unknown field question'
  m.normalize_template_pagination(d,{r._tr for r in t.rows})
  self.assertFalse(list(t._tbl.iter(qn('w:cantSplit'))))
 def composed(self,width=1,text=None,status='answered'):
  tmp=tempfile.TemporaryDirectory();self.addCleanup(tmp.cleanup);root=Path(tmp.name);self.report_root=root
  d=Document(ROOT/'skill/assets/official/edpb-2026-v1-en.docx')
  # Keep the actual template's numbered styles, but use a small isolated table.
  for element in list(d.element.body):
   if element.tag!=qn('w:sectPr'):d.element.body.remove(element)
  table=d.add_table(rows=2,cols=1);table.cell(0,0).text='Retention period';table.cell(1,0).width=Inches(width)
  template=root/'template.docx';d.save(template)
  assessment=root/'assessment.json';assessment.write_text(json.dumps({'meta':{'assessment_id':'test','revision_id':'r1'},'risks':[],'actions':[],'rule_assessments':[],'sources':[]}))
  mapping=root/'map.json';mapping.write_text(json.dumps({'fields':[{'key':'retention','table':0,'row':1,'cell':0,'mode':'blank','label':'Retention period'}]}))
  answer=text or 'Raw records are retained for thirty days. Review and deletion evidence remain unverified. '*4
  field={'key':'retention','status':status,'label':'Retention period','text':answer,'rationale':'Test evidence','factIds':[],'ruleIds':[],'sourceIds':[]}
  if status=='unknown':field.update(question='How long?',impact='Affects retention',recommendation='Confirm policy')
  job={'basisDigest':'basis','files':{'assessment':{'path':str(assessment),'sha256':m.sha(assessment)}},'fieldMap':str(mapping),'templatePath':str(template),'template':{'language':'en','official':True},'templateId':'fixture','language':'en','reportType':'draft'}
  content={'basisDigest':'basis','fields':[field],'risks':[],'dimensions':[],'actions':[]}
  m.compose(job,content,root)
  return Document(root/'report.docx'),json.loads((root/'field-review.json').read_text()),answer
 def test_column_header_is_reused_without_new_label(self):
  d,review,answer=self.composed(text='30 days')
  self.assertEqual(d.tables[0].cell(1,0).text.strip(),'30 days')
  header=d.tables[0].rows[0]
  self.assertIsNotNone(header._tr.find(qn('w:trPr')).find(qn('w:cantSplit')))
  self.assertTrue(header.cells[0].paragraphs[0].paragraph_format.keep_together)
 def test_narrow_answer_is_preserved_in_full_width_note_and_bound(self):
  d,review,answer=self.composed();f=review['fields'][0]
  self.assertNotIn(answer,d.tables[0].cell(1,0).text)
  self.assertIn('F001',d.tables[0].cell(1,0).text)
  ps=list(d.element.body.iter(qn('w:p')))
  actual=[''.join(t.text or '' for t in ps[i].iter(qn('w:t'))) for i in f['answerParagraphs']]
  self.assertEqual(actual,[answer]);self.assertEqual(f['answerSha256'],hashlib.sha256(answer.strip().encode()).hexdigest())
  headings=[p for p in d.paragraphs if p.text.startswith('F001')]
  self.assertEqual(len(headings),1)
  self.assertEqual(headings[0]._p.find(qn('w:pPr')).find(qn('w:numPr')).find(qn('w:numId')).get(qn('w:val')),'0')
 def test_unknown_narrow_answer_stays_blank_with_native_comment(self):
  d,review,_=self.composed(status='unknown');self.assertEqual(d.tables[0].cell(1,0).text.strip(),'Retention period')
  bindings=json.loads((self.report_root/'bindings.json').read_text());field=bindings['unknown_fields'][0]
  ps=list(d.element.body.iter(qn('w:p')));cell=d.tables[0].cell(1,0)._tc
  self.assertIn(ps[field['label_paragraph']],list(cell.iter(qn('w:p'))))
  self.assertTrue(all(not ''.join(t.text or '' for t in ps[i].iter(qn('w:t'))).strip() for i in field['answer_paragraphs']))
  self.assertEqual(len(list(d.comments)),1);self.assertEqual(review['fields'][0]['status'],'unknown')
 def test_wide_answer_is_not_moved(self):
  d,review,answer=self.composed(width=6);self.assertIn(answer,d.tables[0].cell(1,0).text)
  self.assertEqual(review['fields'][0]['paragraphs'],review['fields'][0]['answerParagraphs'])
 def answer_errors(self):
  root=self.report_root
  args={'mode':'docx','assessment':str(root/'assessment.json'),'document':str(root/'report.docx'),'bindings':str(root/'bindings.json'),'review':str(root/'field-review.json'),'language':'en'}
  result=subprocess.run([sys.executable,'-B',str(ROOT/'report-check.py')],input=json.dumps(args),text=True,capture_output=True)
  self.assertEqual(result.returncode,0,result.stdout+result.stderr)
  return [e for e in json.loads(result.stdout)['errors'] if 'authored answer' in e or 'answer binding' in e]
 def test_checker_detects_missing_note_content_after_docx_hash_is_updated(self):
  d,review,_=self.composed();self.assertEqual(self.answer_errors(),[])
  ps=list(d.element.body.iter(qn('w:p')));p=m.Paragraph(ps[review['fields'][0]['answerParagraphs'][0]],d);p.text='Omitted the original answer.'
  d.save(self.report_root/'report.docx')
  bp=self.report_root/'bindings.json';b=json.loads(bp.read_text());b['docx_sha256']=m.sha(self.report_root/'report.docx');bp.write_text(json.dumps(b))
  self.assertEqual(self.answer_errors(),['retention: authored answer changed or omitted'])
 def test_multiline_answer_binding_uses_visible_text(self):
  self.composed(text='Thirty days.\nDeletion verification pending.')
  self.assertEqual(self.answer_errors(),[])
 def test_orphan_ref_preserves_visible_original_and_records_defect(self):
  d=Document(ROOT/'skill/assets/official/edpb-2026-v1-en.docx')
  visible=lambda:''.join(t.text or '' for t in d.element.body.iter(qn('w:t')))
  before=visible();notes=m.preserve_orphan_references(d)
  self.assertEqual(notes,['6.2']);self.assertEqual(visible(),before)
  self.assertFalse(any((t.text or '').lstrip().startswith('REF ') for t in d.element.body.iter(qn('w:instrText'))))
  self.assertIn('not a valid link',list(d.comments)[-1].text)
  self.assertTrue(any('PAGEREF' in (t.text or '') for t in d.element.body.iter(qn('w:instrText'))))
 def test_multiline_append_label_never_duplicates_decision_options(self):
  d=Document();table=d.add_table(rows=1,cols=1);c=table.cell(0,0)
  c.paragraphs[0].add_run('REJECTED: abandon.');c.add_paragraph('APPROVED: proceed.')
  count=len(c.paragraphs);p=m.field_label(c,table.rows[0]._tr,d,'REJECTED: abandon. APPROVED: proceed.','append')
  self.assertEqual(len(c.paragraphs),count);self.assertEqual(p.text,'REJECTED: abandon.')
 def test_actual_composite_label_reuses_original_row_and_header(self):
  d=Document(ROOT/'skill/assets/official/edpb-2026-v1-en.docx');t=d.tables[17];row=t.rows[1];c=row.cells[1]
  before=c.text;original=''.join(t.text or '' for t in d.element.body.iter(qn('w:t')))
  p=m.field_label(c,row._tr,d,'Fairness principle — List of supporting measures','blank',t._tbl)
  self.assertIn('Fairness',p.text);self.assertEqual(c.text,before)
  m.preserve_table_headers(d)
  self.assertEqual(''.join(t.text or '' for t in d.element.body.iter(qn('w:t'))),original)
  for table in d.tables:
   self.assertIsNotNone(table.rows[0]._tr.find(qn('w:trPr')).find(qn('w:cantSplit')))
 def test_short_narrow_risk_answer_moves_without_breaking_reference_or_text(self):
  text='Undetermined pending verification and reassessment.'
  self.assertLess(len(text),80)
  d,review,answer=self.composed(width=.7,text=text)
  p=next(p for p in d.tables[0].cell(1,0).paragraphs if 'F001' in p.text)
  self.assertTrue(p.paragraph_format.keep_together)
  self.assertNotIn(text,d.tables[0].cell(1,0).text)
  self.assertIn(text,[p.text for p in d.paragraphs]);self.assertEqual(self.answer_errors(),[])
 def test_note_identifier_is_one_link_to_its_complete_answer(self):
  d,review,text=self.composed(width=.7,text='Residual likelihood undetermined pending verification.')
  c=d.tables[0].cell(1,0)
  self.assertEqual(c.text.strip(),'F001')
  link=c._tc.find('.//'+qn('w:hyperlink'));self.assertIsNotNone(link)
  anchor=link.get(qn('w:anchor'));targets=[b for b in d.element.body.iter(qn('w:bookmarkStart')) if b.get(qn('w:name'))==anchor]
  self.assertEqual(len(targets),1)
  title=m.Paragraph(targets[0].getparent(),d)
  self.assertTrue(title.text.startswith('F001 —'))
  self.assertEqual(m.Paragraph(title._p.getnext(),d).text,text)
  self.assertEqual(self.answer_errors(),[])
 def test_wider_slash_connected_answer_is_preserved_in_note(self):
  text='Only P1/P2; no sale/advertising/ML/customer analytics; new purpose requires fresh assessment.'
  d,review,answer=self.composed(width=1.9,text=text)
  self.assertEqual(d.tables[0].cell(1,0).text.strip(),'F001')
  self.assertIn(answer,[p.text for p in d.paragraphs]);self.assertEqual(self.answer_errors(),[])
 def test_missing_label_precedes_blank_answer_position(self):
  d=Document();table=d.add_table(rows=1,cols=1);c=table.cell(0,0)
  p=m.field_label(c,table.rows[0]._tr,d,'Decision justification','blank',table._tbl)
  self.assertIs(c.paragraphs[0]._p,p._p);self.assertEqual(c.paragraphs[1].text,'')
if __name__=='__main__':unittest.main()
