"""Compose reviewed structured content into DOCX. Does not invent legal answers."""
import json,sys,hashlib,re,unicodedata
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import _Cell
from docx.text.paragraph import Paragraph

sha=lambda p:hashlib.sha256(Path(p).read_bytes()).hexdigest()
def required(value,label):
    if not isinstance(value,str) or not value.strip():raise ValueError(label+' is required')
    return value.strip()
def choice_options(cell):
    out=[]
    for paragraph in cell.iter(qn('w:p')):
        for symbol in paragraph.iter(qn('w:sym')):
            if symbol.get(qn('w:font'))=='Wingdings 2' and symbol.get(qn('w:char'))=='F02A':
                out.append((symbol,''.join(t.text or '' for t in paragraph.iter(qn('w:t'))).strip()))
    return out
def outline_choices(job):
    if not job.get('fieldMap'):return {}
    doc=Document(job['templatePath']);tables=doc.element.body.findall(qn('w:tbl'));out={}
    for f in json.loads(Path(job['fieldMap']).read_text())['fields']:
        cell=tables[f['table']].findall(qn('w:tr'))[f['row']].findall(qn('w:tc'))[f['cell']]
        choices=choice_options(cell)
        if choices:out[f['key']]=[{'index':i,'label':label} for i,(_,label) in enumerate(choices)]
    return out
def preserve_orphan_references(doc):
    """Keep the publisher's visible wording, but expose its missing targets.

    An orphan REF cannot be recalculated by Word/LibreOffice. Preserve its
    cached text as plain text and attach a provenance comment, never invent a
    replacement section or silently present it as a valid report reference.
    """
    bookmarks={x.get(qn('w:name')) for x in doc.element.body.iter(qn('w:bookmarkStart'))}
    unresolved=[]
    for element in doc.element.body.iter(qn('w:p')):
        runs=list(element.iter(qn('w:r')));start=None;instruction='';separated=False;cached=[]
        for i,run in enumerate(runs):
            for field in run.findall(qn('w:fldChar')):
                kind=field.get(qn('w:fldCharType'))
                if kind=='begin':start=i;instruction='';separated=False;cached=[]
                elif kind=='separate':separated=True
                elif kind=='end' and start is not None:
                    match=re.match(r'\s*REF\s+(\S+)',instruction)
                    if match and match[1] not in bookmarks:
                        text=''.join(cached).strip()
                        if not text:raise ValueError('Orphan template reference has no visible text')
                        for member in runs[start:i+1]:
                            for child in list(member):
                                if child.tag in [qn('w:fldChar'),qn('w:instrText')]:member.remove(child)
                        paragraph=Paragraph(element,doc)
                        doc.add_comment(paragraph.runs,text='Source-template reference '+text+' has no target in the supplied original. Its wording is retained as plain text; it is not a valid link to a section of this assessment.',author='DPIA workbench',initials='DPIA')
                        unresolved.append(text)
                    start=None;separated=False
            instruction+=''.join(x.text or '' for x in run.findall(qn('w:instrText')))
            if separated:cached.extend(x.text or '' for x in run.findall(qn('w:t')))
    return unresolved

def field_label(cell,row,doc,label,mode,table=None):
    normal=lambda value:re.sub(r'\s+','',value)
    candidates=[Paragraph(p,doc) for p in row.iter(qn('w:p'))]
    found=next((p for p in candidates if normal(label) in normal(p.text) and p.runs),None)
    if found is not None:return found
    if mode=='blank' and table is not None:
        header=table.find(qn('w:tr'))
        found=next((Paragraph(p,doc) for p in header.iter(qn('w:p')) if normal(label) in normal(Paragraph(p,doc).text) and Paragraph(p,doc).runs),None)
        if found is not None:return found
        # Field maps combine row and column labels; neither paragraph alone
        # contains the full label. Reuse only when both parts actually exist.
        parts=label.split(' — ',1)
        if len(parts)==2:
            row_label=next((p for p in candidates if normal(parts[0]) in normal(p.text) and p.runs),None)
            column_label=next((Paragraph(p,doc) for p in header.iter(qn('w:p')) if normal(parts[1]) in normal(Paragraph(p,doc).text) and Paragraph(p,doc).runs),None)
            if row_label is not None and column_label is not None:return row_label
    # Append destinations already contain the publisher's question/options.
    # Repeating the concatenated field-map label can look like a second answer.
    if mode=='append':
        found=next((p for p in cell.paragraphs if p.text.strip() and p.runs),None)
        if found is not None:return found
    label_p=cell.add_paragraph(label)
    if mode=='blank':cell.paragraphs[0]._p.addprevious(label_p._p)
    return label_p
def preserve_table_headers(doc):
    # Prevent short column headings from being split across pages. Do not
    # lock whole data rows: a checkbox/answer row can exceed one page.
    for table in doc.element.body.findall(qn('w:tbl')):
        header=table.find(qn('w:tr'))
        if header is None:continue
        props=header.get_or_add_trPr()
        if props.find(qn('w:cantSplit')) is None:props.append(OxmlElement('w:cantSplit'))
        for element in header.iter(qn('w:p')):
            p=Paragraph(element,doc)
            p.paragraph_format.keep_together=True
            p.paragraph_format.keep_with_next=True
def needs_field_note(cell,text):
    if not cell.width:return False
    if cell.width<=Inches(1.65) and len(text)>80:return True
    # Conservative estimate at 10.5 pt, reserving cell margins. Even a short
    # answer can contain a word too wide for a very narrow risk column.
    available=max(0,cell.width/914400*72-14)
    return any(sum(10.5 if unicodedata.east_asian_width(c) in ('W','F') else 6.3 for c in word)>available for word in text.split())
def normalize_template_pagination(doc,note_rows):
    # A page-break-only paragraph can itself overflow to a fresh page before
    # breaking again. Move a single such break onto the following heading.
    # Never remove bookmarks, comments, fields, sections or deliberate double
    # breaks. Work only on plain empty body paragraphs, outside mapped cells.
    def plain_empty(element):
        if element.tag!=qn('w:p'):return False
        for child in element:
            if child.tag==qn('w:pPr'):
                if child.find(qn('w:sectPr')) is not None:return False
            elif child.tag==qn('w:r'):
                if any(x.tag not in [qn('w:rPr'),qn('w:br')] for x in child):return False
                if any(x.get(qn('w:type'))!='page' for x in child.findall(qn('w:br'))):return False
            else:return False
        return True
    for element in list(doc.element.body):
        if element.tag!=qn('w:p'):continue
        heading=Paragraph(element,doc)
        if not heading.style or not heading.style.name.startswith('Heading'):continue
        empty=[];previous=element.getprevious()
        while previous is not None and plain_empty(previous):
            empty.append(previous);previous=previous.getprevious()
        breaks=[b for p in empty for b in p.iter(qn('w:br'))]
        # Other explicit page/section boundaries are not part of this repair.
        if len(breaks)!=1 or any(p.find('./'+qn('w:pPr')+'/'+qn('w:pageBreakBefore')) is not None for p in empty):continue
        heading.paragraph_format.page_break_before=True
        for p in empty:doc.element.body.remove(p)
    for row in note_rows:
        # Only generated compact reference rows qualify. A checkbox, unknown
        # field or long answer must still be allowed to flow across pages.
        cells=row.findall(qn('w:tc'))
        texts=[''.join(t.text or '' for t in c.iter(qn('w:t'))).strip() for c in cells]
        if not all(re.fullmatch(r'(?:F\d{3,}|\d+)?',text) for text in texts):continue
        if not any(re.fullmatch(r'F\d{3,}',text) for text in texts):continue
        if any(len(list(c.iter(qn('w:p'))))>3 for c in cells):continue
        if any(next(row.iter(qn('w:'+tag)),None) is not None for tag in ['drawing','pict','object','sym','fldChar','commentRangeStart','sectPr','trHeight']):continue
        props=row.get_or_add_trPr()
        if props.find(qn('w:cantSplit')) is None:props.append(OxmlElement('w:cantSplit'))
def compose(job,content,out):
    if content.get('basisDigest')!=job['basisDigest']:raise ValueError('Content basis changed')
    assessment=json.loads(Path(job['files']['assessment']['path']).read_text())
    if sha(job['files']['assessment']['path'])!=job['files']['assessment']['sha256']:raise ValueError('Assessment bytes changed')
    mapping=json.loads(Path(job['fieldMap']).read_text())['fields'] if job.get('fieldMap') else []
    fields=content.get('fields',[]);keys=[f['key'] for f in fields]
    if not fields or len(keys)!=len(set(keys)):raise ValueError('Unique nonempty fields required')
    if mapping and set(keys)!={f['key'] for f in mapping}:raise ValueError('Every mapped template field must be authored explicitly')
    for f in fields:
        if f.get('status') not in ['answered','unknown','not_applicable']:raise ValueError('Invalid field status')
        for ref in ['factIds','ruleIds','sourceIds']:
            if not isinstance(f.get(ref),list):raise ValueError(ref+' required')
        required(f.get('rationale'),'Field rationale')
        if f['status']=='unknown':
            for key in ['question','impact','recommendation']:required(f.get(key),key)
            if job['reportType']=='formal' and f.get('category')!='signature':raise ValueError('Formal substantive field cannot be unknown')
        else:required(f.get('text'),'Field answer')
    risks=content.get('risks',[]);dimensions=content.get('dimensions',[]);actions=content.get('actions',[])
    if job['templateId']=='legal-elements' and not {'scope','facts_sources','processing','screening','measures_actions','consultation','review_conditions','decision'}.issubset(keys):raise ValueError('Self-authored report is missing required sections')
    for values,expected,key in [(risks,[r['id'] for r in assessment['risks']],'id'),(actions,[a['id'] for a in assessment['actions']],'id'),(dimensions,[r['rule_id'] for r in assessment['rule_assessments'] if r['rule_id'].startswith('EU-C')],'ruleId')]:
        ids=[r[key] for r in values]
        if len(ids)!=len(set(ids)) or set(ids)!=set(expected):raise ValueError('Risk/dimension coverage differs from assessment')
        for r in values:required(r.get('text'),'Authored risk/dimension text')
    out=Path(out)
    for name in ['report.docx','bindings.json','field-review.json']:
        if (out/name).exists():raise ValueError('Composition output exists; preserve it and use a new report attempt')
    native_pdf=job.get('template',{}).get('format')=='pdf'
    if native_pdf:
        import importlib.util
        spec=importlib.util.spec_from_file_location('report_pdf',Path(__file__).with_name('report-pdf.py'));pdf=importlib.util.module_from_spec(spec);spec.loader.exec_module(pdf)
        pdf.fill(job,content,out)
        job={**job,'templatePath':None,'fieldMap':None}
    doc=Document(job['templatePath']) if job.get('templatePath') else Document()
    # The supplied template contains cached PAGEREF results for its empty form.
    # Retain linked section names, but never publish those stale page numbers.
    for paragraph in doc.element.body.iter(qn('w:p')):
        runs=list(paragraph.iter(qn('w:r')));in_page_ref=False;result=False
        for run in runs:
            for instr in run.findall(qn('w:instrText')):
                if 'PAGEREF ' in (instr.text or ''):in_page_ref=True
            for field in run.findall(qn('w:fldChar')):
                if field.get(qn('w:fldCharType'))=='separate' and in_page_ref:result=True
                elif field.get(qn('w:fldCharType'))=='end' and in_page_ref:in_page_ref=False;result=False
            if result:
                for text in run.findall(qn('w:t')):text.text=''
    orphan_references=preserve_orphan_references(doc)
    native_paragraphs=list(doc.element.body.iter(qn('w:p')))
    original_language=job['template']['language'] if job['template'].get('official') else content.get('templateVerification',{}).get('originalLanguage',job['language'])
    if job.get('templatePath') and original_language in [None,'user_selected']:raise ValueError('Specify the actual original template language')
    translated=bool(job.get('templatePath') and original_language!=job['language']);translations=[]
    if translated:
        translations=content.get('translations',[])
        # Use visible text only; OOXML properties are not content.
        original={i:p for i,p in enumerate(native_paragraphs) if ''.join(t.text or '' for t in p.iter(qn('w:t'))).strip()}
        if len(translations)!=len(original) or {t['paragraph'] for t in translations}!=set(original):raise ValueError('Translation must cover every original nonblank paragraph')
        for item in translations:
            p=Paragraph(original[item['paragraph']],doc);style=p.style;p.clear();p.style=style;p.add_run(required(item.get('text'),'Translated template text'))
    if not job.get('templatePath'):
        normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(10.5);normal.paragraph_format.space_after=Pt(7)
        normal.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'PingFang SC')
        for section in doc.sections:section.top_margin=section.bottom_margin=Inches(.75)
        doc.add_heading(content.get('title') or 'DPIA',0)
    en=job['language']=='en'
    notice=('DRAFT — incomplete assessment; review open comments.' if job['reportType']=='draft' else 'Assessment report — prepared from the confirmed report basis; not regulatory approval.') if en else ('待完善草稿｜请审阅未决字段及原生批注。' if job['reportType']=='draft' else '正式交付版｜依据已确认的报告范围生成，不代表监管批准。')
    if translated:notice+=' 非官方译本 / Unofficial translation.'
    if job.get('fieldMap'):notice+=(' Checkbox key: ☒ selected; ☐ not selected; [?] unresolved. F-number links refer to full answers in Field notes. Contents retain section links; cached template page numbers are omitted.' if en else ' 选项说明：☒ 已选；☐ 未选；[?] 未决。F 编号链接指向字段补充说明中的完整答案。目录保留章节链接，不沿用空模板的缓存页码。')
    if orphan_references:notice+=(' Source-template cross-references '+', '.join(orphan_references)+' have missing targets; their original wording is retained as plain text with comments, not as valid report links.' if en else ' 原模板交叉引用 '+', '.join(orphan_references)+' 缺失目标；原文字样保留为纯文本并加批注，不作为本报告的有效章节链接。')
    if job['template'].get('official'):notice+=(' Publisher attribution/signature retained from the blank template is not endorsement of this assessment.' if en else ' 空模板保留的发布机构署名不代表该机构认可本评估。')
    if native_pdf:notice+=' 转换工作副本 / Converted working copy. 原生模板及填写结果以随附 PDF 为准 / The accompanying PDF preserves the original template.'
    marker=doc.add_paragraph(notice);marker.runs[0].bold=True;doc.element.body.insert(0,marker._p)
    for section in doc.sections:
        # Official templates may keep the footer in a content control/table.
        # Preserve that structure instead of assuming a plain paragraph exists.
        if not section.footer.paragraphs:continue
        footer=section.footer.paragraphs[0]
        if not footer.text.strip():
            footer.alignment=2;footer.add_run('DPIA · ');field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    if job.get('templatePath'):preserve_table_headers(doc)
    map_by_key={f['key']:f for f in mapping};pending=[];field_nodes=[];note_fields=[];answer_nodes={};note_rows=set()
    bookmarks=list(doc.element.body.iter(qn('w:bookmarkStart')))
    bookmark_id=max([int(b.get(qn('w:id'))) for b in bookmarks]+[-1])+1
    bookmark_names={b.get(qn('w:name')) for b in bookmarks}
    for f in fields:
        loc=map_by_key.get(f['key']) or f.get('location')
        if job.get('templatePath'):
            if not loc:raise ValueError('Uploaded template field needs table/row/cell location')
            ti,ri,ci=(loc[k] for k in ['table','row','cell'])
            if any(type(i)!=int or i<0 for i in [ti,ri,ci]):raise ValueError('Invalid template position')
            table=doc.element.body.findall(qn('w:tbl'))[ti];row=table.findall(qn('w:tr'))[ri];cell=_Cell(row.findall(qn('w:tc'))[ci],doc)
            choices=choice_options(cell._tc)
            if choices:
                selections=f.get('selections')
                if selections is None and f['status']=='unknown':selections=[{'index':i,'status':'unknown','reason':f['rationale']} for i in range(len(choices))]
                if not isinstance(selections,list) or len(selections)!=len(choices) or {x.get('index') for x in selections}!=set(range(len(choices))):raise ValueError(f['key']+': every checkbox option needs an explicit selection status')
                for selection in selections:
                    status=selection.get('status');required(selection.get('reason'),'Checkbox reason')
                    if status not in ['selected','not_selected','unknown']:raise ValueError('Invalid checkbox status')
                    if job['reportType']=='formal' and status=='unknown':raise ValueError('Formal report has an unresolved checkbox')
                    symbol,_=choices[selection['index']];run=symbol.getparent();run.remove(symbol)
                    mark=OxmlElement('w:t');mark.text={'selected':'☒','not_selected':'☐','unknown':'[?]'}[status];run.append(mark)
                    props=run.find(qn('w:rPr'))
                    if props is not None:
                        fonts=props.find(qn('w:rFonts'))
                        if fonts is not None:
                            for attr in ['ascii','hAnsi','eastAsia']:fonts.set(qn('w:'+attr),'Arial Unicode MS')
                f['selections']=selections
            label=required(f.get('label') or (map_by_key.get(f['key']) or {}).get('label'),'Field label')
            # Unknown-field comments must remain local to their own row, not a shared header.
            label_p=field_label(cell,row,doc,label,loc.get('mode'),table if f['status']!='unknown' else None)
            # Blank destinations can contain formatting; retain it. Append fields
            # keep the template's original question and option text intact.
            answer_p=next((p for p in cell.paragraphs if not p.text and p._p is not label_p._p),None) if loc.get('mode')=='blank' else None
            if answer_p is None:answer_p=cell.add_paragraph()
        else:
            label_p=doc.add_paragraph(required(f.get('label'),'Field label'));label_p.runs[0].bold=True;label_p.paragraph_format.keep_with_next=True;answer_p=doc.add_paragraph()
        if f['status']=='unknown':
            if not label_p.runs:label_p.add_run(f.get('label') or f['key'])
            comment=doc.add_comment(label_p.runs,text=f['question']+'\n'+f['impact']+'\n'+f['recommendation'],author='DPIA workbench',initials='DPIA')
            pending.append((f,label_p._p,answer_p._p,str(comment.comment_id)))
        else:
            answer_p.alignment=0
            # Keep the original table and questions intact. Long generated
            # answers in narrow cells retain their full text in bound notes.
            if job.get('templatePath') and needs_field_note(cell,f['text']):
                note_id='F'+str(len(note_fields)+1).zfill(3)
                anchor='DPIA_Field_'+note_id
                while anchor in bookmark_names:anchor+='x'
                bookmark_names.add(anchor)
                # A single linked identifier cannot separate its introductory
                # words from its number at a table page break in LibreOffice.
                link=OxmlElement('w:hyperlink');link.set(qn('w:anchor'),anchor)
                run=OxmlElement('w:r');props=OxmlElement('w:rPr')
                for name,value in [('sz','21'),('color','0563C1'),('u','single')]:
                    item=OxmlElement('w:'+name);item.set(qn('w:val'),value);props.append(item)
                run.append(props);text=OxmlElement('w:t');text.text=note_id;run.append(text);link.append(run);answer_p._p.append(link)
                answer_p.paragraph_format.keep_together=True
                note_fields.append((f,note_id,anchor))
                note_rows.add(row)
            else:
                answer_p.add_run(f['text']);answer_nodes[f['key']]=answer_p._p
            for run in answer_p.runs:run.font.size=Pt(10.5)
        field_nodes.append((f,answer_p._p))
    doc.add_heading('Assessment dimensions' if en else '评估维度及限制',1);dimension_nodes=[]
    for item in dimensions:
        doc.add_heading(item['ruleId'],2);p=doc.add_paragraph(item['text']);dimension_nodes.append((item,p._p))
    doc.add_heading('Risk reasoning and evidence' if en else '风险推理与依据',1)
    reference_file=Path(__file__).parent/'skill/references/sources.json';legal={s['id']:s['title'] for s in json.loads(reference_file.read_text())['sources']};legal.update({s['id']:s['title'] for s in assessment['sources']})
    actual={r['id']:r for r in assessment['risks']};risk_nodes=[]
    for item in risks:
        record=actual[item['id']];start=doc.add_heading(item['id'],2);doc.add_paragraph(item['text'])
        rating=doc.add_paragraph(({'high':'high','not_high':'not high','undetermined':'undetermined'} if en else {'high':'高','not_high':'非高','undetermined':'待定'})[record['gdpr_rating']['level']])
        end=doc.add_paragraph(('References: ' if en else '依据：')+'；'.join(legal[ref['source_id']]+' '+ref['locator'] for ref in record['legal_refs']))
        risk_nodes.append((item,start._p,end._p,rating._p))
    doc.add_heading('Actions and reassessment' if en else '整改计划与复评',1);action_nodes=[]
    action_records={a['id']:a for a in assessment['actions']}
    for item in actions:
        record=action_records[item['id']];doc.add_heading(item['id'],2);p=doc.add_paragraph(item['text']);action_nodes.append((item,p._p))
        doc.add_paragraph(('Recorded status: ' if en else '登记状态：')+record['status'])
        for key,title in [('owner','Owner' if en else '责任人'),('due','Due date' if en else '期限')]:
            label_p=doc.add_paragraph(title);label_p.paragraph_format.keep_with_next=True;answer_p=doc.add_paragraph(record[key])
            if not record[key].strip():
                f={'key':'action:'+item['id']+':'+key,'status':'unknown','rationale':'Not recorded in the assessment' if en else '底稿尚未记录','factIds':[],'ruleIds':record['rule_ids'],'sourceIds':record['evidence_source_ids'],
                   'question':('Please identify '+title+' for '+item['id']) if en else '请确定 '+item['id']+' 的'+title,'impact':'Affects action tracking and completion.' if en else '影响整改跟踪与完成判断。','recommendation':'Record the responsible decision and reassess.' if en else '补充实际安排后复评。'}
                comment=doc.add_comment(label_p.runs,text=f['question']+'\n'+f['impact']+'\n'+f['recommendation'],author='DPIA workbench',initials='DPIA');pending.append((f,label_p._p,answer_p._p,str(comment.comment_id)));field_nodes.append((f,answer_p._p))
    if note_fields:
        doc.add_heading('Field notes' if en else '字段补充说明',1)
        for f,note_id,anchor in note_fields:
            # Plain numbered identifiers must not inherit template outline
            # numbering; otherwise labels can render twice (10.1 10.1).
            title=doc.add_paragraph(note_id+' — '+(f.get('label') or map_by_key[f['key']]['label']))
            start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(bookmark_id));start.set(qn('w:name'),anchor)
            end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(bookmark_id));bookmark_id+=1
            title.runs[0]._r.addprevious(start);title._p.append(end)
            title.runs[0].bold=True;title.paragraph_format.keep_with_next=True
            num=OxmlElement('w:numPr');num_id=OxmlElement('w:numId');num_id.set(qn('w:val'),'0');num.append(num_id);title._p.get_or_add_pPr().append(num)
            p=doc.add_paragraph(f['text']);p.alignment=0;answer_nodes[f['key']]=p._p
    for element in doc.element.body.iter(qn('w:p')):
        p=Paragraph(element,doc)
        if p.style and p.style.name.startswith('Heading'):p.paragraph_format.keep_with_next=True
    if job.get('templatePath'):normalize_template_pagination(doc,note_rows)
    ps=list(doc.element.body.iter(qn('w:p')));positions={p:i for i,p in enumerate(ps)}
    path=out/'report.docx';doc.save(path)
    bindings={'assessment_sha256':sha(job['files']['assessment']['path']),'docx_sha256':sha(path),'assessment_id':assessment['meta']['assessment_id'],'revision_id':assessment['meta']['revision_id'],
              'risks':[{'id':r['id'],'start':positions[start],'end':positions[end],'rating_paragraph':positions[rating]} for r,start,end,rating in risk_nodes],
              'unknown_fields':[{'id':f['key'],'label_paragraph':positions[label],'answer_paragraphs':[positions[answer]],'comment_id':cid,**{k:f[k] for k in ['question','impact','recommendation']}} for f,label,answer,cid in pending]}
    review={'basisDigest':job['basisDigest'],'templateVerification':content.get('templateVerification'),
            'fields':[{k:v for k,v in f.items() if k not in ['text','location']}|{'paragraphs':list(dict.fromkeys([positions[p]]+([positions[answer_nodes[f['key']]]] if f['key'] in answer_nodes else [])))}|({'answerParagraphs':[positions[answer_nodes[f['key']]]],'answerSha256':hashlib.sha256(re.sub(r'\s+',' ',f['text']).strip().encode()).hexdigest()} if f['key'] in answer_nodes else {}) for f,p in field_nodes],
            'dimensions':[{k:v for k,v in item.items() if k!='text'}|{'paragraphs':[positions[p]]} for item,p in dimension_nodes],
            'actions':[{'id':item['id'],'paragraphs':[positions[p]]} for item,p in action_nodes],
            'translationBindings':[{'sourceParagraph':item['paragraph'],'paragraph':positions[native_paragraphs[item['paragraph']]],'text':item['text'],'preserveReason':item.get('preserveReason','')} for item in translations],
            'translationReview':content.get('translationReview'),
            'nativePdf':{'bindingsSha256':sha(out/'pdf-bindings.json'),'format':'pdf','wordRole':'converted_working_copy'} if native_pdf else None,
            'pages':[],'contentReview':content.get('contentReview',{'status':'pending','notes':'Await authored content review'})}
    for name,data in [('bindings.json',bindings),('field-review.json',review)]:
        with (out/name).open('x') as f:json.dump(data,f,ensure_ascii=False,indent=2)
    return {'document':str(path),'unknownFields':len(pending),'risks':len(risks),'dimensions':len(dimensions),'actions':len(actions),'pageReview':'not_started'}
if __name__=='__main__':
    try:
        args=json.load(sys.stdin);print(json.dumps(outline_choices(args['job']) if args.get('mode')=='choices' else compose(args['job'],json.loads(Path(args['content']).read_text()),args['out']),ensure_ascii=False))
    except Exception as e:print(json.dumps({'error':str(e)[:2000]}));sys.exit(1)
